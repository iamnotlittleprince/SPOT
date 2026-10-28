from pathlib import Path
import re, json, hashlib
from datetime import datetime, timezone
from pptx import Presentation
from pptx.util import Inches, Pt

root=Path(__file__).resolve().parents[1]
source=Path(__file__).resolve().parent/'fonte-corrigida-convertida.pptx'
prs=Presentation(source)
assert len(prs.slides)==11

def replace_text(shape,text):
    tf=shape.text_frame
    first=tf.paragraphs[0]
    if first.runs:
        first.runs[0].text=text
        for run in first.runs[1:]:run.text=''
    else:first.add_run().text=text
    for p in list(tf.paragraphs)[1:]:tf._txBody.remove(p._p)

def find(slide,text):
    return next(sh for sh in slide.shapes if sh.has_text_frame and sh.text.strip()==text)

cover=prs.slides[0]
replace_text(find(cover,'participantes'),'integrantes')
replace_text(find(cover,'15'),'11')
replace_text(find(cover,'15 mi'),'15 min')
replace_text(find(cover,'roteiro estimado'),'tempo máximo')
prs.core_properties.title='SPOT — Defesa de TCC — versão corrigida e revisada'
prs.core_properties.subject='11 slides; revisão visual de 28/10/2026; fonte: recuperação com alterações do usuário'
prs.core_properties.last_modified_by='Equipe SPOT'
prs.core_properties.modified=datetime(2026,10,28,12,0,tzinfo=timezone.utc)

# Numeração acompanha a ordem escolhida pelo usuário.
for i,s in enumerate(prs.slides):
    for sh in s.shapes:
        if not sh.has_text_frame:continue
        t=sh.text.strip()
        if re.fullmatch(r'\d+ / \d+',t):replace_text(sh,f'{i+1:02d} / 11')
        elif i and re.match(r'^\d+ / [^\d]',t):
            title=t.split(' / ',1)[1]
            if i==1:title='PROPOSTA'
            replace_text(sh,f'{i:02d} / {title}')

# Arquitetura: ordem real de comunicação, com alinhamento uniforme.
s=prs.slides[8]
for title,desc,x in [('React','Telas e formulários',.74),('API Laravel','Sessão, acesso e validação',3.94),('Domínio + Eloquent','Regras e transações',7.14),('PostgreSQL','Dados e relacionamentos',10.34)]:
    head=find(s,title);body=find(s,desc)
    head.left=body.left=Inches(x)
    head.top=Inches(2.56);head.width=Inches(2.27)
    head.height=Inches(.57 if title=='Domínio + Eloquent' else .36)
    body.top=Inches(3.23);body.width=Inches(2.27);body.height=Inches(.30)
    for p in head.text_frame.paragraphs:
        for run in p.runs:run.font.size=Pt(16)

# Requisitos: normaliza as duas linhas que foram reordenadas manualmente.
s=prs.slides[5]
for label,percent,count,feature,y in [
 ('Organização ajuda no faturamento','88,3%','83/94','Horas + custos + resultado',2.22),
 ('Dificuldade de organização','63,8%','60/94','Projetos + tarefas + filtros',3.28),
 ('Dificuldade de alinhamento','41,5%','39/94','Equipes + responsáveis',4.22),
 ('Dificuldade de custos adicionais','31,9%','30/94','Despesas + aprovação',5.16),
]:
    h=find(s,label);h.left=Inches(.60);h.top=Inches(y);h.width=Inches(4.10);h.height=Inches(.62)
    h=find(s,percent);h.left=Inches(8.14);h.top=Inches(y-.03);h.width=Inches(1.48);h.height=Inches(.39)
    h=find(s,count);h.left=Inches(8.23);h.top=Inches(y+.42);h.width=Inches(1);h.height=Inches(.22)
    h=find(s,feature);h.left=Inches(9.65);h.top=Inches(y+.06);h.width=Inches(3.05);h.height=Inches(.65)
# As barras mantêm os valores escolhidos; acompanham a nova grade vertical.
bar_shapes=[sh for sh in s.shapes if abs(sh.left/Inches(1)-4.8)<.02 and sh.width/Inches(1)>1 and sh.height/Inches(1)<.25]
for sh in bar_shapes:
    old_y=sh.top/Inches(1)
    if old_y<3:sh.top=Inches(2.28)
    elif old_y<4:sh.top=Inches(3.34)
    elif old_y<5:sh.top=Inches(4.28)
    else:sh.top=Inches(5.22)

# Converte as notas antigas para a ordem atual; não reintroduz os slides removidos.
times=[20,40,50,60,65,70,65,65,60,60,75]
rows=[];elapsed=0
for i,(s,duration) in enumerate(zip(prs.slides,times)):
    notes=s.notes_slide.notes_text_frame
    raw=notes.text
    raw=re.sub(r'^SLIDE \d+[^\n]*\nAcumulado:[^\n]*\n*','',raw).strip()
    raw=raw.replace('<número>','').strip()
    if i==1:
        raw=raw.replace('Obrigado.','Esses são os três pilares que orientam a proposta a seguir.')
    if i==10:
        source_pos=raw.find('Fontes:')
        closing='Para concluir: o Spot conecta os requisitos levantados com regras verificáveis de organização, governança e resultado gerencial. Obrigado.\n\n'
        raw=raw[:source_pos]+closing+raw[source_pos:] if source_pos>=0 else raw+'\n\n'+closing
    # Fala de comparação acompanha o recorte de seis sistemas mantido no slide.
    if i==6:
        raw=raw.replace(', além das ferramentas brasileiras e de automação citadas no trabalho','')
        raw=raw.replace(' TaskRush também trabalha com margem, então rentabilidade não é invenção nossa. AWS é infraestrutura, não concorrente direto.',' Recursos semelhantes existem em outras ferramentas; defendemos a adequação ao cliente, não exclusividade.')
    start=f'{elapsed//60:02d}:{elapsed%60:02d}';elapsed+=duration;end=f'{elapsed//60:02d}:{elapsed%60:02d}'
    notes.text=f'SLIDE {i+1} — TEMPO: {duration} segundos\nAcumulado: {start} a {end}\n\n{raw}'
    rows.append({'slide':i+1,'segundos':duration,'inicio':start,'fim':end})

out=root/'SPOT-Defesa-TCC-revisado.pptx'
prs.save(out)
report={'data_revisao':'2026-10-28','fonte_recuperacao':'/home/alexander/.config/libreoffice/4/user/backup/SPOT-Defesa-TCC-15min.pptx_1.odp','fonte_sha256':hashlib.sha256((Path(__file__).parent/'fonte-corrigida-recuperada.odp').read_bytes()).hexdigest(),'slides':11,'roteiro_segundos':elapsed,'tempo_maximo_segundos':900,'ordem_usuario_preservada':True,'originais_preservados':True,'tempos':rows}
(Path(__file__).parent/'revisao.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({'saida':str(out),'slides':11,'roteiro_segundos':elapsed},ensure_ascii=False))
