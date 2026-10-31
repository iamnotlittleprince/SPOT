from pathlib import Path
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'docs/tcc/SPOT-Defesa-TCC-revisado.pptx'
p=Presentation(Path(__file__).parent/'antes-identidade/SPOT-Defesa-TCC-revisado.pptx')
colors={'007e98':'0C588A','00869c':'0C588A','008b79':'0C588A','64d9e6':'E31B36','112c42':'123044','183447':'123044','6b58a8':'0C588A'}
for s in p.slides:
 for el in s._element.iter():
  if el.tag.endswith('}srgbClr') and el.get('val','').lower() in colors:el.set('val',colors[el.get('val').lower()])
 for sh in s.shapes:
  if sh.has_text_frame:
   for par in sh.text_frame.paragraphs:
    for r in par.runs:r.font.name='Aptos'

def text(s,t,x,y,w,h,size=18,color='123044',bold=False):
 sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for j,line in enumerate(t.split('\n')):
  par=tf.paragraphs[0] if j==0 else tf.add_paragraph();par.text=line
  for r in par.runs:r.font.name='Aptos';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 return sh

def rect(s,x,y,w,h,color):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background();return sh
s=p.slides.add_slide(p.slide_layouts[6])
for placeholder in list(s.placeholders):
 placeholder._element.getparent().remove(placeholder._element)
s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('FFFFFF')
text(s,'11 / PRÓXIMOS PASSOS',.6,.32,9,.3,12,'0C588A',True)
text(s,'O que falta para entrar em produção',.6,.8,12,.6,30,bold=True)
text(s,'Etapa final: implantação, autenticação e validação com o cliente.',.6,1.56,12,.5,17,'526878')
items=[('01','Implantar no Azure','Publicar o Spot em uma máquina virtual Linux no Azure.'),('02','Ativar as autenticações','Configurar Google e Microsoft no ambiente de produção.'),('03','Executar testes de segurança','Rodar os testes de segurança e corrigir os pontos identificados.'),('04','Receber o feedback do cliente','Validar o uso com o cliente e incorporar os ajustes necessários.')]
for i,(num,title,desc) in enumerate(items):
 x=.6+(i%2)*6.15;y=2.35+(i//2)*2.05
 rect(s,x,y,5.95,1.75,'F3F6F8');rect(s,x,y,.055,1.75,'0C588A')
 text(s,num,x+.22,y+.2,.55,.4,19,'0C588A',True)
 text(s,title,x+.88,y+.22,4.8,.68,19,bold=True)
 text(s,desc,x+.88,y+1.0,4.65,.65,17,'526878')
text(s,'Pendências para produção • planejamento da equipe SPOT',.6,7.12,10,.2,9,'526878')
s.notes_slide.notes_text_frame.text='Para finalizar, restam quatro frentes: publicar o Spot em uma VM Linux no Azure, configurar as autenticações Google e Microsoft em produção, executar os testes de segurança e receber o feedback do cliente para os ajustes finais. Essas atividades são pendências, não etapas já concluídas.'
# Marca corporativa no cabeçalho, com recorte nativo do espaço transparente do PNG.
logo=ROOT/'public/computecnica-logo.png'
im=Image.open(logo).convert('RGBA');bbox=im.getchannel('A').getbbox();W,H=im.size
for idx,s in enumerate(p.slides):
 rect(s,0,7.46,13.333,.04,'E31B36')
 if idx==0:
  rect(s,9.5,6.35,3.1,.65,'FFFFFF')
  cover_logo=s.shapes.add_picture(str(logo),Inches(9.7),Inches(6.47),width=Inches(2.7),height=Inches(.4))
  if bbox:
   l,t,r,b=bbox;cover_logo.crop_left=l/W;cover_logo.crop_right=1-r/W;cover_logo.crop_top=t/H;cover_logo.crop_bottom=1-b/H
 if idx>0:
  pic=s.shapes.add_picture(str(logo),Inches(11.1),Inches(.19),width=Inches(1.65),height=Inches(.42))
  if bbox:
   l,t,r,b=bbox;pic.crop_left=l/W;pic.crop_right=1-r/W;pic.crop_top=t/H;pic.crop_bottom=1-b/H
  pic._element.nvPicPr.cNvPr.set('descr','Computécnica')
  text(s,'SPOT',10.15,.3,.8,.28,14,'0C588A',True)
 for sh in s.shapes:
  if idx>0 and sh.has_text_frame and re.match(r'^\d+ / [^\d]',sh.text.strip()):
   for par in sh.text_frame.paragraphs:
    for r in par.runs:
     r.text=re.sub(r'^\d+ / ',f'{idx:02d} / ',r.text)
  if sh.has_text_frame and re.fullmatch(r'\d+ / \d+',sh.text.strip()):
   for par in sh.text_frame.paragraphs:
    for r in par.runs:r.text=''
   sh.text_frame.paragraphs[0].text=f'{idx+1:02d} / 12'
   for r in sh.text_frame.paragraphs[0].runs:r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string('526878')
# Contagem da capa.
for sh in p.slides[0].shapes:
 if sh.has_text_frame and sh.text.strip()=='11':
  for par in sh.text_frame.paragraphs:
   for r in par.runs:
    if r.text.strip()=='11':r.text='12'
text(s,'12 / 12',12,7.09,.75,.25,10,'526878')
p.core_properties.subject='12 slides — identidade Computécnica e Spot; pendências para produção'
p.save(OUT)
print(OUT)
