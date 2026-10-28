from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
p=Path(__file__).resolve().parents[1]/'SPOT-Defesa-TCC-revisado.pptx'
r=Presentation(p)
NAVY='123044';TEAL='00869C';GRAY='526D80';BG='F3F6F7';WHITE='FFFFFF'
def box(s,x,y,w,h,color,round=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background();return sh

def txt(s,x,y,w,h,t,size=18,color=NAVY,bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(t.split('\n')):
  para=tf.paragraphs[0] if i==0 else tf.add_paragraph();para.text=line;para.font.name='Aptos';para.font.size=Pt(size);para.font.bold=bold;para.font.color.rgb=RGBColor.from_string(color);para.space_after=Pt(5)
 return sh

def base(index,kicker,title,sub):
 s=r.slides[index]
 for sh in list(s.shapes):s.shapes._spTree.remove(sh._element)
 s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
 box(s,0,0,13.333,.07,TEAL)
 txt(s,.6,.32,12,.25,kicker,11,TEAL,True);txt(s,.6,.80,12.1,.7,title,28,NAVY,True);txt(s,.6,1.56,12,.5,sub,14,GRAY)
 txt(s,.6,7.12,11,.2,'Fonte: código local do Spot • revisão 28/10/2026',8,GRAY)
 txt(s,12.1,7.09,.65,.25,f'{index+1:02d} / 11',10,GRAY)
 return s

def banner(s,t):
 box(s,.6,6.46,12.12,.49,NAVY,True);txt(s,.78,6.56,11.75,.3,t,12,WHITE,True)
s=base(8,'08 / DIFERENCIAIS','Diferenciais que conectam gestão e resultado','O diferencial está na combinação e na adequação à Computécnica; recursos isolados existem no mercado.')
cards=[('Financeiro ligado à execução','Horas, tarifas e despesas compõem o resultado gerencial do projeto.'),('Regras para a operação do cliente','Contratos, permissões e encerramento com justificativa no mesmo fluxo.'),('Rastreabilidade das alterações','Histórico de ações e captura da tarifa válida na data trabalhada.'),('Portfólio para decidir','Indicadores, filtros, agrupamentos, gráficos e exportação CSV.'),('Integrações e rotina','Google/Microsoft OAuth, agenda, Teams, convites e lembretes.'),('Controle sobre a evolução','Código próprio permite adaptar regras e escolher a hospedagem.')]
for i,(title,body) in enumerate(cards):
 x=.6+(i%3)*4.12;y=2.23+(i//3)*1.95
 box(s,x,y,3.88,1.73,WHITE,True);box(s,x,y,.045,1.73,TEAL)
 txt(s,x+.18,y+.17,3.5,.55,title,19,TEAL,True);txt(s,x+.18,y+.80,3.5,.77,body,16)
banner(s,'Valor da proposta: organização + governança + visão financeira em um fluxo adaptável ao cliente.')
s.notes_slide.notes_text_frame.text='''SLIDE 9 — TEMPO: 60 segundos
Acumulado: 07:15 a 08:15

O diferencial do Spot está na combinação de recursos adaptada à Computécnica. A execução das tarefas alimenta a visão financeira; regras de acesso e encerramento dão controle à operação; a auditoria e a tarifa capturada permitem rastrear o que ocorreu. O portfólio reúne indicadores para decisão. As integrações dependem da configuração das credenciais e os envios por e-mail exigem transporte e worker ativos. Código próprio permite adequação e hospedagem escolhida pelo cliente, mas não elimina custos de manutenção. Não alegamos exclusividade de funções que também existem nos concorrentes.'''
s=base(9,'09 / FLUXO TÉCNICO','Da interação ao dado, ao indicador e ao aviso','Núcleo integrado em Laravel; interface React e processos de apoio com responsabilidades definidas.')
steps=[('1  React','Telas, formulários e filtros\nEnvia a requisição'),('2  API Laravel','Sessão e permissões\nValida os dados'),('3  Domínio + Eloquent','Regras e transações\nTarifas, custos e auditoria'),('4  PostgreSQL','Persiste os registros\nRelaciona os dados')]
for i,(title,body) in enumerate(steps):
 x=.6+i*3.1;box(s,x,2.22,2.80,1.35,NAVY,True);txt(s,x+.15,2.38,2.5,.4,title,14 if i==2 else 17,WHITE,True);txt(s,x+.15,2.89,2.5,.57,body,13,'DCE8EE')
 if i<3:txt(s,x+2.84,2.72,.26,.4,'→',19,TEAL,True)
box(s,.6,3.76,12.12,.43,'E3EDF0',True);txt(s,.77,3.85,11.8,.26,'Retorno: PostgreSQL → Laravel → JSON → React atualiza a tela e os indicadores.',13,NAVY,True)
for i,(title,body) in enumerate([('Indicadores • Python + Pandas','Laravel envia JSON → Pandas agrega → devolve JSON. O script não acessa o banco diretamente.'),('E-mails • fila + worker','Convite/ativação → fila → validação do link → transporte SMTP configurado; retentativas.'),('Rotinas • scheduler','Agenda lembretes e notificações; aplica retenção de auditoria e limpeza de mensagens.')]):
 x=.6+i*4.12;box(s,x,4.43,3.88,1.72,WHITE,True);box(s,x,4.43,.045,1.72,TEAL);txt(s,x+.17,4.60,3.52,.55,title,17,TEAL,True);txt(s,x+.17,5.18,3.52,.85,body,14)
banner(s,'Monólito organizado: regras no backend, persistência central e tarefas de apoio em processos separados.')
s.notes_slide.notes_text_frame.text='''SLIDE 10 — TEMPO: 60 segundos
Acumulado: 08:15 a 09:15

O usuário interage com o React, que envia uma requisição para a API Laravel. O backend identifica a sessão, verifica permissões e valida os campos. As regras de negócio usam o Eloquent para consultar e persistir dados no PostgreSQL, com transações nas operações que precisam ser atômicas. A resposta volta em JSON e a interface atualiza sua apresentação. Para o portfólio, o Laravel seleciona os dados e executa o Python com Pandas, que recebe e devolve JSON, sem consultar o banco. Convites e ativações entram na fila: o worker verifica se o registro ainda é utilizável antes de enviar pelo transporte configurado e repete tentativas em caso de falha. O scheduler agenda lembretes e processamento de notificações, retenção de auditoria e limpeza de mensagens. Essa separação de processos não transforma o sistema em microsserviços. O próximo slide detalha uma tarefa nesse fluxo.'''
r.save(p)
print(p)
