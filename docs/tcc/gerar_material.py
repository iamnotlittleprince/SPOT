#!/usr/bin/env python3
"""Gera slides, gráficos e guia a partir das fontes auditadas em dados_pesquisa.json.
Dependências: python-pptx, python-docx, matplotlib, Pillow; PyMuPDF para verificação.
Executar da raiz: PYTHONPATH=/tmp/spot-tcc-libs python3 docs/tcc/gerar_material.py
"""
from pathlib import Path
import json, csv, re, textwrap
from PIL import ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from docx import Document
from docx.shared import Cm, Pt as DPt, RGBColor as DColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT=Path(__file__).resolve().parent
A=ROOT/'assets'; G=A/'graficos';G.mkdir(exist_ok=True)
data=json.loads((ROOT/'dados_pesquisa.json').read_text())
Q={q['id']:q for q in data['questoes']}
NAVY='112C42';CYAN='007E98';TEAL='008B79';AMBER='AF6714';RED='B34B4B';INK='183447';MUTED='526878';BG='F3F6F8';LINE='DCE5EA';WHITE='FFFFFF';PURPLE='6B58A8'
with (ROOT/'dados_pesquisa.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['figura','questao','alternativa','contagem','percentual_base94','soma_figura','diferenca_nao_discriminada','nota'])
 for q in data['questoes']:
  for lab,v in zip(q['rotulos'],q['valores']):
   w.writerow([q['figura'],q['titulo'],lab,v,round(v/94*100,2),sum(q['valores']),'' if q.get('multipla') else 94-sum(q['valores']),q.get('nota','')])
 for q in data['qualitativa']['categorias']:
  w.writerow(['Tabela 1','Sugestões abertas',q['tema'],q['frequencia'],'','','','Frequências temáticas publicadas; 49 na tabela vs. 57 comentários no texto.'])

# Gráficos de apoio exportáveis; todos com a procedência e a base explícitas.
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False})
for q in data['questoes']+[{'id':'qualitativa','titulo':'Sugestões abertas — frequências publicadas','rotulos':[x['tema'] for x in data['qualitativa']['categorias']],'valores':[x['frequencia'] for x in data['qualitativa']['categorias']],'qual':True}]:
 fig,ax=plt.subplots(figsize=(10,4.9),layout='constrained')
 vals=q['valores'];labs=q['rotulos'];colors=['#'+CYAN]*len(vals)
 ax.barh(range(len(vals)),vals,color=colors,height=.58);ax.set_yticks(range(len(vals)),labs);ax.invert_yaxis();ax.set_xlim(0,max(vals)*1.26)
 ax.set_title(q['titulo'],loc='left',fontweight='bold',pad=18,color='#'+NAVY)
 for i,v in enumerate(vals):ax.text(v+.3,i,str(v),va='center',fontweight='bold',color='#'+INK)
 ax.set_xlabel('Frequência temática' if q.get('qual') else ('Menções (não exclusivas)' if q.get('multipla') else 'Participantes nas categorias exibidas'))
 ax.xaxis.grid(True,alpha=.12);ax.set_axisbelow(True)
 note='Fonte: pesquisa SPOT (2026), tabela 1. Tabela = 49 frequências; texto = 57 comentários.' if q.get('qual') else f"Fonte: pesquisa SPOT (2026), figura {q['figura']}. Base declarada: 94. "
 if not q.get('qual'):
  note+=('Múltiplas menções; não somar percentuais.' if q.get('multipla') else f"Soma exibida: {sum(vals)}; diferença não discriminada: {94-sum(vals)}.")
 fig.text(.02,-.035,note,fontsize=8,color='#'+MUTED)
 fig.savefig(G/f"{q['id']}.png",dpi=180,bbox_inches='tight');fig.savefig(G/f"{q['id']}.svg",bbox_inches='tight');plt.close(fig)

prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
prs.core_properties.title='SPOT — Da pesquisa à implementação';prs.core_properties.author='Equipe SPOT';prs.core_properties.subject='Defesa de TCC — 15 slides, 14min10s'
slides=[]
def color(v):return RGBColor.from_string(v)
def rect(s,x,y,w,h,fill=WHITE,line=None,round=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=color(fill);sh.line.fill.background() if not line else None
 if line:sh.line.color.rgb=color(line)
 if round:
  try:sh.adjustments[0]=.08
  except Exception:pass
 return sh

def txt(s,text,x,y,w,h,size=18,c=INK,bold=False,font='DejaVu Sans',align=None):
 box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=Inches(.015);tf.margin_top=tf.margin_bottom=0
 for j,line in enumerate(str(text).split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(c);p.space_after=Pt(5)
  if align:p.alignment=align
 return box

def base(k,title,sub='',source='Fonte: código local do Spot • corte técnico 26/10/2026'):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color(BG)
 rect(s,0,0,13.333,.08,CYAN);txt(s,k.upper(),.55,.32,11,.25,10,CYAN,True)
 txt(s,title,.55,.77,12.2,.72,29,NAVY,True)
 if sub:txt(s,sub,.57,1.49,12.1,.53,13,MUTED)
 txt(s,source,.55,7.12,11.5,.22,8,MUTED);txt(s,f'{len(slides)+1:02d} / 15',12.05,7.09,.73,.25,10,MUTED,align=PP_ALIGN.RIGHT)
 slides.append(s);return s

def text_height(text,w,size,bold=False):
 font_path='/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf' if bold else '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
 font=ImageFont.truetype(font_path,round(size*4))
 limit=(w-.06)*72*4*.96
 lines=0
 paragraphs=str(text).split('\n')
 for para in paragraphs:
  current='';count=1
  for word in para.split():
   candidate=f'{current} {word}'.strip()
   if current and font.getlength(candidate)>limit:
    count+=1;current=word
   else:current=candidate
  lines+=count
 return max(.26,(lines*size*1.22+(len(paragraphs)-1)*5)/72+.045)

def card(s,title,body,x,y,w,h,accent=CYAN):
 rect(s,x,y,w,h,WHITE,LINE,True);rect(s,x,y,.055,h,accent)
 title_h=text_height(title,w-.35,18,True)
 body_y=y+.15+title_h+.17
 txt(s,title,x+.18,y+.15,w-.35,title_h,18,accent,True)
 txt(s,body,x+.18,body_y,w-.35,max(.3,y+h-.16-body_y),15)

def side_notes(s,items,x,y,w,bottom=6.65):
 # A altura do título inclui as quebras antes de posicionar sua descrição.
 for title,body in items:
  title_h=text_height(title,w,18,True)
  body_h=text_height(body,w,15)
  txt(s,title,x,y,w,title_h,18,CYAN,True)
  body_y=y+title_h+.16
  txt(s,body,x,body_y,w,body_h,15)
  y=body_y+body_h+.28
 assert y-.28<=bottom+.02, f'Coluna de notas excede área disponível: {y-.28:.2f}'

def banner(s,text,y=6.42,fill=NAVY):
 rect(s,.55,y,12.22,.48,fill,round=True);txt(s,text,.72,y+.1,11.9,.27,12,WHITE,True)

def mini(s,id,x,y,w,h,fs=12,short=None):
 q=Q[id];rect(s,x,y,w,h,WHITE,LINE,True)
 title_h=text_height(q['titulo'],w-.3,14,True)
 txt(s,q['titulo'],x+.16,y+.12,w-.3,title_h,14,NAVY,True)
 content_top=.12+title_h+.18
 labels=short or q['rotulos'];vals=q['valores'];labelw=w*.47;barw=w-labelw-1.17
 rowh=(h-content_top-.39)/len(vals);m=max(vals)
 for i,(label,v) in enumerate(zip(labels,vals)):
  yy=y+content_top+i*rowh
  txt(s,label,x+.15,yy,labelw-.12,max(.19,rowh-.015),fs,MUTED)
  rect(s,x+labelw,yy+.035,barw,.14,LINE)
  rect(s,x+labelw,yy+.035,barw*v/m,.14,CYAN)
  txt(s,f'{v} · {v/94*100:.1f}%'.replace('.',','),x+labelw+barw+.06,yy,1.04,.25,max(8,fs-2),INK,True)
 tail='menções não exclusivas • % de 94' if q.get('multipla') else f"soma: {sum(vals)} • % sobre base 94"
 txt(s,tail,x+.16,y+h-.25,w-.3,.18,8,MUTED)

def age_line(s,x,y,w,h):
 q=Q['idade'];rect(s,x,y,w,h,WHITE,LINE,True)
 txt(s,'Perfil etário — gráfico de linhas',x+.14,y+.12,w-.28,.4,14,NAVY,True)
 cd=CategoryChartData();cd.categories=['15–25','25–35','35–45','45–55','>55']
 cd.add_series('Participantes (%)',[v/94 for v in q['valores']])
 chart=s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS,Inches(x+.12),Inches(y+.63),Inches(w-.24),Inches(h-1.15),cd).chart
 chart.has_legend=False;chart.has_title=False
 chart.chart_style=13
 axis=chart.value_axis;axis.minimum_scale=0;axis.maximum_scale=.4;axis.major_unit=.1
 axis.tick_labels.number_format='0%';axis.tick_labels.font.size=Pt(9)
 axis.tick_labels.font.name='DejaVu Sans';axis.has_major_gridlines=True
 axis.major_gridlines.format.line.color.rgb=color(LINE)
 chart.category_axis.tick_labels.font.size=Pt(9);chart.category_axis.tick_labels.font.name='DejaVu Sans'
 chart.category_axis.major_tick_mark=XL_TICK_MARK.NONE
 series=chart.series[0];series.format.line.color.rgb=color(CYAN);series.format.line.width=Pt(2.5)
 series.marker.style=XL_MARKER_STYLE.CIRCLE;series.marker.size=7
 series.marker.format.fill.solid();series.marker.format.fill.fore_color.rgb=color(CYAN)
 series.marker.format.line.color.rgb=color(WHITE)
 plot=chart.plots[0];plot.has_data_labels=True
 plot.data_labels.position=XL_LABEL_POSITION.ABOVE
 plot.data_labels.font.size=Pt(10);plot.data_labels.font.bold=True;plot.data_labels.font.color.rgb=color(INK)
 for point,value in zip(series.points,q['valores']):
  point.data_label.position=XL_LABEL_POSITION.ABOVE
  point.data_label.text_frame.text=f'{value/94*100:.1f}%'.replace('.',',')
  for paragraph in point.data_label.text_frame.paragraphs:
   paragraph.font.name='DejaVu Sans';paragraph.font.size=Pt(10);paragraph.font.bold=True;paragraph.font.color.rgb=color(INK)
   for run in paragraph.runs:
    run.font.name='DejaVu Sans';run.font.size=Pt(10);run.font.bold=True;run.font.color.rgb=color(INK)
 txt(s,'Base: 94 • linha entre faixas etárias, não evolução no tempo.',x+.14,y+h-.35,w-.28,.28,8,MUTED)

def need_pie(s,x,y,w,h):
 rect(s,x,y,w,h,WHITE,LINE,True)
 txt(s,'Software dedicado — necessidade',x+.14,y+.12,w-.28,.3,13,NAVY,True)
 labels=['Concordam','Talvez','Não','Diferença documental']
 counts=[43,33,14,4];palette=[TEAL,CYAN,AMBER,'A5B5BF']
 cd=CategoryChartData();cd.categories=labels;cd.add_series('Participantes',counts)
 chart=s.shapes.add_chart(XL_CHART_TYPE.PIE,Inches(x+.05),Inches(y+.51),Inches(1.43),Inches(1.34),cd).chart
 chart.has_legend=False;chart.has_title=False
 plot=chart.plots[0];plot.first_slice_angle=270
 for point,co in zip(chart.series[0].points,palette):
  point.format.fill.solid();point.format.fill.fore_color.rgb=color(co);point.format.line.color.rgb=color(WHITE)
 for i,(lab,n,co) in enumerate(zip(labels,counts,palette)):
  yy=y+.56+i*.29;rect(s,x+1.54,yy+.025,.10,.10,co)
  txt(s,f'{lab}: {n/94*100:.1f}% ({n})'.replace('.',','),x+1.70,yy,w-1.81,.28,9,INK,i==0)
 txt(s,'Base 94; diferença de 4 não discriminada na fonte.',x+.14,y+h-.23,w-.28,.18,7.5,MUTED)

def screenshot(s,path,x,y,w,h):
 from PIL import Image
 im=Image.open(path);iw,ih=im.size;ratio=min(w/iw,h/ih);ww=iw*ratio;hh=ih*ratio
 rect(s,x-.025,y-.025,w+.05,h+.05,WHITE,LINE,True)
 s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))

def pill(s,label,x,y,w,fill=CYAN):
 rect(s,x,y,w,.35,fill,round=True);txt(s,label,x+.06,y+.065,w-.12,.21,10,WHITE,True,align=PP_ALIGN.CENTER)

# 01
s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color(NAVY);slides.append(s)
rect(s,9.25,0,4.083,7.5,CYAN)
txt(s,'TCC • INFORMÁTICA • 2026',.7,.6,8,.35,13,'64D9E6',True)
txt(s,'SPOT',.65,1.55,8,1.25,70,WHITE,True)
txt(s,'Da pesquisa\nà implementação',.72,2.95,8.2,1.65,36,WHITE,True)
txt(s,'Projetos, pessoas e resultado financeiro\nno contexto da Computécnica',.75,4.93,7.7,.85,19,'C6DCE5')
txt(s,'Alexander Molany Longo • Lucas Rodrigues Soares\nPedro Galvão de Matos • Yasmin Regina de Souza',.75,6.25,8,.6,11,'C6DCE5')
for i,(n,l) in enumerate([('94','participantes'),('15','slides'),('14:10','roteiro estimado')]):
 txt(s,n,9.68,1.3+i*1.75,3.1,.75,44,WHITE,True);txt(s,l,9.73,2.12+i*1.75,3,.4,15,WHITE)
#02
s=base('01 / ponto de partida','Um problema real. Duas pesquisas.','Como transformar a rotina do cliente em requisitos verificáveis.','Fontes: monografia, introdução; capítulo da pesquisa, metodologia (2026).')
card(s,'Computécnica','Visual antigo\nAcesso pouco delimitado\nCustos desconectados\nDependência da solução anterior',.55,2.22,3.88,3.7)
card(s,'Pesquisa de campo','94 respostas\n18 mar → 24 ago de 2026\nSurvey quali-quantitativo\nAmostra de conveniência',4.72,2.22,3.88,3.7)
card(s,'Pesquisa de sistemas','6 plataformas principais\nIntegrações e automação\nReferências de interação\nEscolhas para o cliente',8.89,2.22,3.88,3.7)
banner(s,'A amostra orienta prioridades; não representa estatisticamente todas as empresas.')
#03
s=base('02 / pesquisa de campo','Quem respondeu e o que já utiliza','88 de 94 estavam empregados • 65 não atuam diretamente em gestão de projetos','Fonte: pesquisa SPOT (2026), figuras 1–3 e perfil. Contagens reconstruídas dos gráficos do DOCX.')
age_line(s,.55,2.18,3.85,3.84)
mini(s,'atuacao',4.74,2.18,3.85,3.84,12,['Não atua','Atua parcialmente','Totalmente','Sim'])
mini(s,'ferramentas',8.93,2.18,3.85,3.84,11,['Office 365','Software próprio','Office 2021','Nenhum','Bloco de Notas','Notion'])
banner(s,'52 menções a Office 365 → integração corporativa • perfil diverso → prioridade de usabilidade')
#04
s=base('03 / dores → requisitos','Organização e financeiro orientam o núcleo','Concordância direta = “Sim” + “Absolutamente” • percentuais sobre 94 participantes','Fonte: pesquisa SPOT (2026), figuras 9–12. Totais exibidos variam; diferenças constam no guia.')
rows=[('Dificuldade de organização',60,'63,8%','Projetos + tarefas + filtros',CYAN),('Organização ajuda no faturamento',83,'88,3%','Horas + custos + resultado',TEAL),('Dificuldade de alinhamento',39,'41,5%','Equipes + responsáveis',PURPLE),('Dificuldade de custos adicionais',30,'31,9%','Despesas + aprovação',AMBER)]
for i,(lab,v,pct,feature,co) in enumerate(rows):
 yy=2.22+i*.94;txt(s,lab,.6,yy,4.1,.35,17,INK,True);rect(s,4.8,yy+.06,3.2,.22,LINE);rect(s,4.8,yy+.06,3.2*v/94,.22,co)
 txt(s,pct,8.14,yy-.03,1.48,.39,22,co,True);txt(s,f'{v}/94',8.23,yy+.41,1,.25,11,MUTED);txt(s,feature,9.65,yy+.06,3.05,.65,16)
banner(s,'Percepções justificam requisitos; não comprovam ganho de faturamento causado pelo Spot.')
#05 all six graphs
s=base('04 / adoção e contexto','Uma solução dedicada precisa se adaptar à rotina','43 concordam com a necessidade; 33 dizem “talvez” → 80,9% de apoio ou abertura, sem inferir compra.','Fonte: pesquisa SPOT (2026), figuras 4–8 e 13. Base 94; alternativas abreviadas; ver transcrição no guia.')
mini(s,'software',.55,2.1,3.87,2.08,10,['Não possui','Possui','Sem foco em projetos'])
need_pie(s,4.72,2.1,3.87,2.08)
mini(s,'satisfacao',8.89,2.1,3.87,2.08,9,['Pode melhorar','Não temos','Atende','Não é usado'])
mini(s,'interface',.55,4.38,3.87,2.08,10,['Sem software','Sim','Mais ou menos','Não','Absolutamente'])
mini(s,'externos',4.72,4.38,3.87,2.08,10)
mini(s,'internos',8.89,4.38,3.87,2.08,10)
txt(s,'Resposta: organização dedicada, flexibilidade de uso e base de modelagem para contextos internos/externos.',.6,6.67,12.1,.28,12,CYAN,True)
#06
s=base('05 / sugestões abertas','Do pedido à entrega — com escopo explícito','57 comentários declarados • tabela publicada: 49 frequências • temas não são depoimentos literais','Fonte: pesquisa SPOT (2026), tabela 1; código atual. Frequências não equivalem a uma distribuição completa.')
entries=[('Interface simples',12,'Tutorial, temas e navegação','Implementado',TEAL),('Financeiro',6,'Horas, tarifas, despesas e resultado','Implementado',TEAL),('Integrações/e-mail',6,'OAuth, agenda e e-mails de atividade','Implementado / configurável',TEAL),('Dashboards',5,'Portfólio, CSV e atualização automática','Implementado',TEAL),('Tarefas/agenda',5,'Visões, agenda e lembretes automáticos','Implementado',TEAL),('Sem sugestão',12,'Sem requisito específico derivado','Categoria da pesquisa',MUTED)]
for i,(theme,num,feature,status,co) in enumerate(entries):
 yy=2.17+i*.59
 rect(s,.55,yy,12.22,.45,WHITE if i%2==0 else 'EAF0F3',round=True)
 txt(s,theme,.7,yy+.08,2.4,.25,13,INK,True);txt(s,str(num),3.1,yy+.08,.5,.25,14,CYAN,True)
 txt(s,feature,3.78,yy+.08,5.2,.27,13);txt(s,status,9.45,yy+.08,3.1,.27,12,co,True)
banner(s,'Recorte funcional da pesquisa • painel atualiza a cada 30 s • e-mail externo exige SMTP e ativação.')
#07
s=base('06 / pesquisa de sistemas','Referências do mercado. Recorte da Computécnica.','Diferencial defendido: adequação das regras e autonomia de evolução, sem alegar exclusividade.','Fontes: monografia, estudo de sistemas; páginas oficiais consultadas em 26/10/2026 (URLs no guia e notas).')
for x,w,t in [(.55,2.0,'Referência'),(2.7,4.2,'Força observada'),(7.22,5.55,'Recorte do Spot')]:
 rect(s,x,2.09,w,.43,NAVY);txt(s,t,x+.1,2.18,w-.2,.23,12,WHITE,True)
bench=[('Monday.com','Gestão visual e automações','Regras financeiras para o cliente'),('Asana','Coordenação de tarefas e projetos','Contrato → horas → resultado'),('Jira','Planejamento e acompanhamento','Operação de serviços e custos'),('Trello','Organização visual em quadros','Tarifas, despesas e governança no domínio'),('ClickUp','Centralização, tempo e dashboards','Modelo específico em código próprio'),('Notion','Documentos e bases configuráveis','Regras transacionais explícitas')]
for i,(a,b,c) in enumerate(bench):
 yy=2.62+i*.43;rect(s,.55,yy,12.22,.39,WHITE if i%2==0 else 'EAF0F3')
 txt(s,a,.68,yy+.07,1.9,.24,13,CYAN,True);txt(s,b,2.8,yy+.07,4.25,.24,13);txt(s,c,7.32,yy+.07,5.25,.24,13)
txt(s,'TaskRush também mede margem • Artia e Jestor exploram automação • AWS é infraestrutura.',.65,5.52,12,.5,14,MUTED)
banner(s,'Código próprio permite adaptação e hospedagem própria; não elimina custo de operação ou manutenção.')
#08
s=base('07 / diferencial financeiro','Do contrato à margem, com histórico preservado','Exemplo didático: contrato de R$ 10.000 + R$ 200 em horas extras faturáveis.')
items=[('Receita gerencial','R$ 10.200',CYAN),('Impostos','− R$ 1.020',AMBER),('Comissão','− R$ 510',AMBER),('Mão de obra','− R$ 350',AMBER),('Despesa aprovada','− R$ 300',AMBER),('Lucro gerencial','R$ 8.020',TEAL)]
for i,(lab,v,co) in enumerate(items):
 xx=.55+(i%3)*4.17;yy=2.15+(i//3)*1.32;rect(s,xx,yy,3.87,1.12,WHITE,LINE,True)
 txt(s,lab,xx+.17,yy+.14,3.5,.29,13,MUTED);txt(s,v,xx+.17,yy+.53,3.5,.43,27,co,True)
for i,(a,b) in enumerate([('Snapshot da tarifa','Reajustes não reescrevem custos capturados.'),('Resultado parcial','Sem tarifa, a pendência fica visível.'),('Cálculo decimal','Regras de impostos e comissão explícitas.')]):
 xx=.65+i*4.17;txt(s,a,xx,5.12,3.8,.35,17,NAVY,True);txt(s,b,xx,5.63,3.7,.55,13)
banner(s,'Receita contratual gerencial ≠ recebimento bancário. Não é um ERP contábil/fiscal completo.')
#09
s=base('08 / experiência de uso','Um fluxo visível para quem executa','Home e tarefas conectam projeto, responsável, participantes, duração e histórico.','Fonte: interface real do build local; dados demonstrativos; API simulada para as capturas.')
screenshot(s,A/'tela-tarefas.png',.55,2.12,8.53,4.48)
side_notes(s,[('Mesmos registros','Quadro • lista • calendário'),('Uso cotidiano','Busca, filtros, prioridade e prazo'),('Aprendizagem','Tutorial • 3 idiomas • temas')],9.42,2.4,3.2)
#10
s=base('09 / portfólio','Indicadores para acompanhar e decidir','Carteira, planejado × realizado, custo por analista/atividade e exportação CSV.','Fonte: interface real do build local; dados demonstrativos; código analítico Laravel + Python/Pandas.')
screenshot(s,A/'tela-portfolio.png',.55,2.1,8.53,4.5)
side_notes(s,[('Personalização','Indicadores, agrupamentos e gráficos'),('Atualização automática','A cada 30 s, após mudanças e entre abas'),('Leitura correta','Período filtra projetos; não é extrato de caixa')],9.42,2.25,3.2)
#11
s=base('10 / governança','Acesso, rastreabilidade e integração no fluxo','Recursos existentes têm regras, dependências e limites definidos.')
card(s,'Governança','Perfis e permissões\nCampos protegidos na API\nFinalização com justificativa\nHistórico e retenção',.55,2.2,3.87,3.85)
card(s,'Integração efetiva','Google / Microsoft OAuth\nCalendários e Teams\nDiretório Microsoft Graph\nRequer credenciais/consentimento',4.72,2.2,3.87,3.85)
card(s,'Convites e rotina','Link de uso único e validade\nTokens protegidos\nE-mails em fila\nWorker e agendador',8.89,2.2,3.87,3.85)
banner(s,'24 atalhos de aplicativos não são 24 integrações de dados. Reabrir projeto não reativa a equipe automaticamente.')
#12
s=base('11 / arquitetura','Um monólito organizado por responsabilidades','Interface e backend no mesmo repositório; processos de apoio com papéis separados.')
nodes=[(.6,'React','Telas e formulários'),(3.8,'API Laravel','Sessão, acesso e validação'),(7.0,'Domínio + Eloquent','Regras e transações'),(10.2,'PostgreSQL','Dados e relacionamentos')]
for i,(xx,a,b) in enumerate(nodes):
 rect(s,xx,2.35,2.55,1.35,NAVY,round=True);txt(s,a,xx+.14,2.57,2.27,.5,17,WHITE,True);txt(s,b,xx+.14,3.12,2.27,.42,11,'CEE4ED')
 if i<3:txt(s,'→',xx+2.63,2.7,.52,.6,30,CYAN,True)
for i,(a,b) in enumerate([('Python + Pandas','JSON → agregação → JSON\nSem conexão direta ao banco'),('Worker de e-mail','Fila → validação do link → SMTP\nRetentativas em segundo plano'),('Scheduler','Retenção de auditoria\nLimpeza periódica de mensagens')]):
 card(s,a,b,.6+i*4.17,4.27,3.82,1.66)
banner(s,'Separar processos não torna a aplicação microserviços; a escolha reduz complexidade de implantação.')
#13
s=base('12 / código em funcionamento','Uma tarefa: do clique à persistência','Exemplo real de caminho no código; regras não ficam apenas na interface.')
steps=[('1','TasksPage.jsx','Horas + minutos → duration_minutes'),('2','api.js → /api/v1/tasks','Sessão + cookie/validação CSRF'),('3','TaskV1Controller.store','Validação, projeto e pessoas autorizadas'),('4','TaskCosting + Eloquent','Tarifa vigente, tarefa e participantes'),('5','AuditRecorder → JSON','Histórico na transação e retorno à tela')]
for i,(n,a,b) in enumerate(steps):
 yy=2.12+i*.78;rect(s,.6,yy,.48,.48,CYAN,round=True);txt(s,n,.64,yy+.07,.4,.3,17,WHITE,True,align=PP_ALIGN.CENTER)
 txt(s,a,1.32,yy+.02,5.4,.34,19,NAVY,True);txt(s,b,7.0,yy+.04,5.65,.5,15)
banner(s,'ProjectFinancialSummary usa snapshots e evita contar a tarefa novamente quando há work_log vinculado.')
#14
s=base('13 / evidências e evolução','O que foi verificado nesta versão local','Validação executada em 26/10/2026; a monografia precisa refletir esse estado técnico.')
card(s,'122 testes aprovados','788 asserções\nBackend em schema de teste\nBuild de produção concluído\n1.583 módulos transformados',.55,2.17,3.87,3.87,TEAL)
card(s,'Entregue no código','Projetos, tarefas e equipes\nPainéis com atualização automática\nAlertas, resumos e silêncio\nIntegrações e e-mails configuráveis',4.72,2.17,3.87,3.87)
card(s,'Operação e validação','Credenciais e consentimentos\nSMTP e envio de atividade habilitado\nValidação com usuários/provedores\nEnsaios de carga e backups',8.89,2.17,3.87,3.87,AMBER)
banner(s,'Testes aprovados ≠ homologação completa. Banco atual: PostgreSQL; referências antigas a SQLite devem mudar.')
#15
s=base('14 / contribuição','Pesquisa → requisitos → implementação','Uma solução adaptada à Computécnica, com caminho técnico verificável.')
for i,(n,a,b) in enumerate([('01','Organizar a execução','Projetos, tarefas, responsáveis e equipes.'),('02','Entender o resultado','Horas, custos, despesas e margem gerencial.'),('03','Controlar a evolução','Permissões, histórico e implantação própria.')]):
 xx=.6+i*4.17;txt(s,n,xx,2.4,3.7,.8,44,CYAN,True);txt(s,a,xx,3.44,3.7,.75,23,NAVY,True);txt(s,b,xx,4.43,3.7,.85,17)
banner(s,'Próximo passo: validar o uso real, medir resultados e evoluir com o retorno do cliente.',6.15)

# Notas: roteiro completo e fontes específicas.
guide=(ROOT/'GUIA-DEFESA-TCC-SPOT.md').read_text()
section=guide.split('## 16. Roteiro oral')[1].split('## 17.')[0]
blocks=re.split(r'### Slide \d+ — ',section)[1:]
times=[20,50,60,70,65,70,75,65,60,55,60,60,55,50,35]
assert len(slides)==15 and len(blocks)==15 and sum(times)==850
acc=0;time_rows=[]
for i,(s,block,sec) in enumerate(zip(slides,blocks,times)):
 heading,speech=block.split('\n',1)
 source='Fontes: pesquisa de campo DOCX (2026), texto e figuras; monografia integrada; código local em 26/10/2026. Consulte o guia para contagens completas e limitações.'
 if i==6:source+=' Páginas oficiais: monday.com/w/use-cases/ai-project-management-software; asana.com/features; atlassian.com/software/jira/features; trello.com/tour; clickup.com/features; notion.com/product/projects; taskrush.com.br/gerenciamento-de-projetos/; artia.com/produto/ia-do-artia/; jestor.com/pt/ai/; aws.amazon.com/what-is-cloud-computing/ (consulta 26/10/2026).'
 s.notes_slide.notes_text_frame.text=f'SLIDE {i+1} — TEMPO: {sec} segundos\nAcumulado: {acc//60:02}:{acc%60:02} a {(acc+sec)//60:02}:{(acc+sec)%60:02}\n\n{speech.strip()}\n\n{source}'
 time_rows.append({'slide':i+1,'titulo':heading.rsplit(' — ',1)[0],'segundos':sec,'inicio':f'{acc//60:02}:{acc%60:02}','fim':f'{(acc+sec)//60:02}:{(acc+sec)%60:02}'})
 acc+=sec
prs.save(ROOT/'SPOT-Defesa-TCC-15min.pptx')
(ROOT/'roteiro-tempos.json').write_text(json.dumps(time_rows,ensure_ascii=False,indent=2))
(ROOT/'ROTEIRO-14min10s.txt').write_text('SPOT — Roteiro oral e tempo\nTempo planejado: 14min10s + margem de 50s até 15min.\n\n'+'\n\n'.join(f"{r['inicio']}–{r['fim']} | Slide {r['slide']} | {r['titulo']}\n{blocks[i].split(chr(10),1)[1].strip()}" for i,r in enumerate(time_rows)))

# Inventário atual de rotas, sem transformar números do framework em funcionalidades.
routes=json.loads((A/'rotas.json').read_text()); app_routes=[r for r in routes if not r.get('vendor') and r.get('uri') not in ['up'] and not str(r.get('action','')).startswith(('Laravel\\','Illuminate\\'))]
appendix=['\n\n## Anexo A. Inventário atual de rotas','',f'Inventário obtido por `./php artisan route:list --json`: {len(routes)} entradas totais (incluindo framework). Subconjunto por critério automático de action fora de Laravel/Illuminate e sem healthcheck: {len(app_routes)}. Esses números não são contagens de funcionalidades.','', '| Método | URI | Ação |','|---|---|---|']
for r in app_routes:appendix.append(f"| {r['method'].replace('|',' / ')} | /{r['uri']} | {r['action'].replace('App'+chr(92)+'Http'+chr(92)+'Controllers'+chr(92),'')} |")
appendix+=['','## Anexo B. Conferência numérica de cada figura','','| Figura | Tema | Soma exibida | Diferença até 94 |','|---|---|---:|---:|']
for q in data['questoes']:appendix.append(f"| {q['figura']} | {q['titulo']} | {sum(q['valores'])} | {'Múltiplas menções' if q.get('multipla') else 94-sum(q['valores'])} |")
appendix+=['','## Anexo C. Gráficos completos da pesquisa','Todos os gráficos abaixo foram reconstruídos dos valores publicados. Nas barras, o valor é contagem, não percentual. A diferença até a base declarada não recebeu categoria inventada.']
for q in data['questoes']:appendix+=['',f"### Figura {q['figura']} — {q['titulo']}",f"![{q['titulo']}](assets/graficos/{q['id']}.png)"]
appendix+=['','### Tabela temática representada em gráfico','![Frequências temáticas](assets/graficos/qualitativa.png)']
full=guide+'\n'.join(appendix)+'\n'
(ROOT/'GUIA-COMPLETO-TCC-SPOT.md').write_text(full)
# TXT com estrutura legível e referências; imagens permanecem referenciadas por caminho.
txtguide=re.sub(r'!\[([^]]+)\]\(([^)]+)\)',r'[Gráfico: \1 | arquivo: \2]',full)
txtguide=re.sub(r'\*\*([^*]+)\*\*',r'\1',txtguide);txtguide=txtguide.replace('`','');txtguide=re.sub(r'^#{1,6} ', '',txtguide,flags=re.M)
(ROOT/'GUIA-COMPLETO-TCC-SPOT.txt').write_text(txtguide)

# Versão A4 para leitura/estudo; títulos e referências acadêmicas, sem alterar a monografia.
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.left_margin=Cm(3);sec.bottom_margin=sec.right_margin=Cm(2)
style=doc.styles['Normal'];style.font.name='Arial';style.font.size=DPt(11);style.paragraph_format.line_spacing=1.3;style.paragraph_format.space_after=DPt(6)
for level in [1,2,3]:
 st=doc.styles[f'Heading {level}'];st.font.name='Arial';st.font.color.rgb=DColor.from_string(NAVY);st.font.size=DPt(16 if level==1 else 13 if level==2 else 11);st.paragraph_format.keep_with_next=True
foot=sec.footer.paragraphs[0];foot.alignment=2;run=foot.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.append(fld)
lines=full.splitlines();in_table=False;in_code=False;table=None
for line in lines:
 if line.startswith('```'):in_code=not in_code;continue
 if not line.strip():continue
 if line.startswith('|') and not in_code:
  cells=[c.strip() for c in line.strip().strip('|').split('|')]
  if all(re.match(r'^:?-+:?$',c) for c in cells):continue
  if not in_table:
   table=doc.add_table(rows=0,cols=len(cells));table.style='Light Shading Accent 1';in_table=True
  row=table.add_row().cells
  for j,c in enumerate(cells[:len(row)]):
   row[j].text=c.replace('`','')
   for p in row[j].paragraphs:
    for r in p.runs:r.font.size=DPt(8)
  continue
 in_table=False
 if line.startswith('!['):
  target=re.search(r'\]\((.*?)\)',line).group(1);doc.add_picture(str(ROOT/target),width=Cm(15.6));continue
 if line.startswith('# '):
  doc.add_heading(line[2:],0);continue
 m=re.match(r'^(#{2,4}) (.*)',line)
 if m:
  if m.group(1)=='##' and m.group(2).startswith(('16.','Anexo C')):doc.add_page_break()
  doc.add_heading(m.group(2),len(m.group(1))-1);continue
 clean=line.replace('**','').replace('`','').rstrip(' ')
 p=doc.add_paragraph(clean)
 if in_code:
  p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=DPt(0)
  for r in p.runs:r.font.name='DejaVu Sans Mono';r.font.size=DPt(7)
 else:p.alignment=3
# Notas pequenas nas tabelas e propriedades do documento.
doc.core_properties.title='SPOT — Guia completo para defesa de TCC';doc.core_properties.author='Equipe SPOT'
doc.save(ROOT/'GUIA-COMPLETO-TCC-SPOT.docx')
print(json.dumps({'slides':len(slides),'tempo_segundos':sum(times),'graficos':len(list(G.glob('*.png'))),'palavras_guia':len(txtguide.split()),'rotas_totais':len(routes),'rotas_criterio_aplicacao':len(app_routes)},ensure_ascii=False))
